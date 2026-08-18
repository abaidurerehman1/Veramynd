"""Export L1 judge vs QA recommended-label match/miss xlsx."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "veramynd_parser" / "prompts" / "L1_QA_Review_for_Engineer.xlsx"
JUDGE = ROOT / "output" / "judge" / "G1M2U1L1.json"
OUT = ROOT / "output" / "reports" / "L1_judge_vs_recommended_match.xlsx"


def main() -> None:
    wb_src = load_workbook(SRC, data_only=True)
    ws_src = wb_src["L1 QA Pass"]
    xlsx_map: dict[str, dict[str, object]] = {}
    for r in range(14, 220):
        code = ws_src.cell(r, 2).value
        rec = ws_src.cell(r, 6).value
        why = ws_src.cell(r, 7).value
        if code is None:
            continue
        code_s = str(code).strip()
        if not re.match(r"^1\.[A-Z]\.", code_s):
            continue
        rec_s = "" if rec is None else str(rec).strip()
        why_s = "" if why is None else str(why).strip()
        labs = re.findall(r"\b(full|partial|none)\b", rec_s.lower())
        xlsx_map[code_s] = {"recommended": rec_s, "labels": labs, "why": why_s}

    j = json.loads(JUDGE.read_text(encoding="utf-8"))
    verdicts = j.get("verdicts") or []
    pred = {v.get("standard_code", "").strip(): v for v in verdicts}

    wb = Workbook()
    ws = wb.active
    ws.title = "L1 Match vs Recommended"
    headers = [
        "Standard",
        "JSON label",
        "Recommended label",
        "Result",
        "In xlsx",
        "In JSON",
        "Notes",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(1, c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")

    match_fill = PatternFill("solid", fgColor="C6EFCE")
    miss_fill = PatternFill("solid", fgColor="FFC7CE")
    match_count = 0

    def add_row(
        code: str,
        json_lbl: str,
        rec_lbl: str,
        result: str,
        in_xlsx: str,
        in_json: str,
        notes: str = "",
    ) -> None:
        nonlocal match_count
        ws.append([code, json_lbl, rec_lbl, result, in_xlsx, in_json, notes])
        r = ws.max_row
        fill = match_fill if result == "MATCH" else miss_fill
        ws.cell(r, 4).fill = fill
        if result == "MATCH":
            match_count += 1

    for v in verdicts:
        code = v.get("standard_code", "").strip()
        cur = (v.get("matched_status") or "").strip().lower()
        if code in xlsx_map:
            info = xlsx_map[code]
            rec = str(info["recommended"])
            labs = info["labels"]
            if not labs:
                result = "MISS"
                notes = "No parsable recommended label"
            elif cur in labs:
                result = "MATCH"
                notes = ""
            else:
                result = "MISS"
                notes = str(info.get("why", ""))[:500]
            add_row(code, cur, rec, result, "Yes", "Yes", notes)
        else:
            add_row(code, cur, "", "MISS", "No", "Yes", "Not in QA xlsx")

    for code in sorted(xlsx_map):
        if code in pred:
            continue
        info = xlsx_map[code]
        add_row(
            code,
            "",
            str(info["recommended"]),
            "MISS",
            "Yes",
            "No",
            "Not in current judge JSON",
        )

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Metric", "Value"])
    ws2.append(["JSON standards", len(verdicts)])
    ws2.append(["QA xlsx standards", len(xlsx_map)])
    ws2.append(["MATCH rows", match_count])
    ws2.append(["Total comparison rows", ws.max_row - 1])
    ws2.append(["MATCH rate on JSON 25", f"{match_count}/{len(verdicts)}"])

    for col in ws.columns:
        letter = col[0].column_letter
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 60)
    ws.column_dimensions["G"].width = 80

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
